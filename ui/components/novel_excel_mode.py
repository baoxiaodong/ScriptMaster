"""
小说Excel分镜生成组件
"""
import io
import logging
import re
import time

import streamlit as st

from core.processor import NovelModeProcessor
from ui.components.error_renderer import show_inline_error
from ui.components.file_upload import render_file_uploader, render_data_preview, render_chapter_selector
from ui.components.results_display import render_results
from utils.error_handler import error_handler
from utils.file_handler import FileHandler
from utils.state_manager import StateManager

logger = logging.getLogger("ScriptMaster.NovelMode")


def render_novel_excel_mode(llm_service, file_handler: FileHandler):
    """主渲染逻辑"""
    st.markdown("### 📚 小说分镜生成")

    # 1. 核心状态初始化 (严格使用 session_state)
    if 'novel_is_generating' not in st.session_state:
        st.session_state.novel_is_generating = False
    if 'novel_results' not in st.session_state:
        st.session_state.novel_results = {}
    if 'novel_outline' not in st.session_state:
        st.session_state.novel_outline = None
    if 'novel_selected_df' not in st.session_state:
        st.session_state.novel_selected_df = None

    is_gen = st.session_state.novel_is_generating
    has_outline = bool(st.session_state.novel_outline)
    has_results = bool(st.session_state.novel_results)

    try:
        # ================= 📌 永远置顶：文件上传框 =================
        uploaded_file = render_file_uploader()
        if uploaded_file is None:
            return

        df = file_handler.read_file(uploaded_file)
        if df is None or df.empty:
            return

        total_episodes = st.session_state.get('total_episodes', 20)

        # ================= 🕳️ 魔法黑洞容器 =================
        main_ui = st.empty()

        # ================= 🚀 大模型生成执行区 (防残影拦截器) =================
        if is_gen:
            main_ui.empty()  # 💥 瞬间炸毁黑洞容器，绝不留残影！

            st.divider()
            st.warning("⏳ 引擎高速运转中，请勿刷新浏览器页面...")

            # 提取被锁定的 DataFrame（如果没有，兜底用全量）
            target_df = st.session_state.get('novel_selected_df')
            if target_df is None or target_df.empty:
                target_df = df

            if not has_outline:
                _execute_generation_flow(llm_service, target_df, total_episodes)
                return
            elif has_outline and not has_results:
                render_novel_outline_section(llm_service, target_df, total_episodes)
                st.markdown("### 📜 生成分镜脚本")
                _execute_scripts_generation(llm_service, target_df, total_episodes)
                return
            elif has_results:
                error_keys = [k for k, v in _sort_and_clean_results(st.session_state.novel_results).items()
                              if (isinstance(v, str) and v.startswith("❌")) or (hasattr(v, 'empty') and v.empty) or (
                                          hasattr(v, '__len__') and len(v) < 20)]
                if error_keys:
                    st.markdown("### 📜 补全分镜脚本")
                    _execute_retry_flow(llm_service, target_df, error_keys,
                                        _sort_and_clean_results(st.session_state.novel_results))
                return

                # ================= 🎨 交互界面渲染区 =================
        with main_ui.container():
            # 【状态 0：初始页面】
            if not has_outline and not has_results:
                render_data_preview(df)
                selected_indices = render_chapter_selector(df)
                # 动态获取用户选中的行
                current_selected_df = df.iloc[selected_indices] if selected_indices else df

                st.divider()
                st.markdown("**⚙️ 生成配置**")
                col1, col2 = st.columns([3, 1])
                with col1:
                    total_episodes = st.number_input(
                        "总集数", min_value=10, max_value=100,
                        value=st.session_state.get('total_episodes', 20), step=5
                    )
                    st.session_state.total_episodes = total_episodes
                with col2:
                    batch_count = (total_episodes + 2) // 3
                    st.markdown(
                        f"""<div style="background-color: #FFF9F0; border: 1px solid #F0D9A0; padding: 10px; border-radius: 8px; text-align: center;">
                            <p style="margin: 0; font-size: 14px; color: #E3700D; font-weight: bold;">
                                分 {batch_count} 次生成<br><span style="color: #C68A1E; font-size: 11px; font-weight: normal;">多线程加速</span>
                            </p>
                        </div>""", unsafe_allow_html=True
                    )

                api_ready = st.session_state.get("api_validated", False)
                if not api_ready:
                    st.warning("⚠️ 请在侧边栏验证 API 配置")
                    st.button("🚀 开始生成大纲", disabled=True)
                else:
                    if st.button(f"🚀 第一步：开始生成 {total_episodes} 集大纲", type="primary",
                                 use_container_width=True):
                        # 🚨 在这里死死锁定用户选中的数据，防止生成时被覆盖！
                        st.session_state.novel_selected_df = current_selected_df
                        st.session_state.novel_is_generating = True
                        st.rerun()

            # 【状态 1：大纲完成】
            elif has_outline and not has_results:
                if st.button("🗑️ 返回初始配置 (清空大纲)", use_container_width=True):
                    st.session_state.novel_outline = None
                    st.session_state.novel_results = {}
                    st.session_state.novel_selected_df = None
                    st.rerun()

                render_novel_outline_section(llm_service, st.session_state.novel_selected_df, total_episodes)

            # 【状态 2：分镜完成】
            elif has_results:
                if st.button("🗑️ 清空所有进度，开始新项目", use_container_width=True):
                    st.session_state.novel_outline = None
                    st.session_state.novel_results = {}
                    st.session_state.novel_selected_df = None
                    st.rerun()

                render_novel_outline_section(llm_service, st.session_state.novel_selected_df, total_episodes)

                sorted_results = _sort_and_clean_results(st.session_state.novel_results)
                render_results(sorted_results)

                error_keys = [k for k, v in sorted_results.items()
                              if (isinstance(v, str) and v.startswith("❌")) or (hasattr(v, 'empty') and v.empty) or (
                                          hasattr(v, '__len__') and len(v) < 20)]

                retry_disabled = not error_keys
                retry_label = "🔄 补全缺失集数" if error_keys else "✅ 全部生成成功，完整无缺"

                if st.button(retry_label, use_container_width=True, disabled=retry_disabled):
                    st.session_state.novel_is_generating = True
                    st.rerun()

    except Exception as e:
        error_handler.handle_exception(exception=e, context="NovelExcelMode", show_page=True)


def render_novel_outline_section(llm_service, df, total_episodes: int = 20):
    """渲染小说大纲展示区域"""
    if not st.session_state.novel_outline:
        return

    if st.session_state.novel_outline.startswith("❌"):
        st.error(st.session_state.novel_outline)
        if st.button("🔄 重新尝试生成大纲", type="primary", use_container_width=True):
            st.session_state.novel_outline = None
            st.session_state.novel_is_generating = True
            st.rerun()
        return

    st.markdown("### 📖 分集大纲节点")
    base_name = st.session_state.get('uploaded_file_name', '小说')

    is_gen = st.session_state.novel_is_generating
    has_results = bool(st.session_state.novel_results)
    is_readonly = is_gen or has_results
    is_expanded = not is_readonly

    with st.expander("📋 查看/编辑剧情大纲", expanded=is_expanded):
        edited_outline = st.text_area(
            "大纲内容（修改后请保存）：",
            value=st.session_state.novel_outline,
            height=350 if is_expanded else 150,  # 折叠时减小高度
            label_visibility="collapsed",
            disabled=is_readonly
        )

        if not is_gen:
            try:
                from docx import Document
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                doc = Document()
                title = doc.add_heading(f'{base_name}_分集大纲', 0)
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                clean_text = '\n'.join([line.strip() for line in
                                        st.session_state.novel_outline.replace('###', '').replace('**', '').replace('*',
                                                                                                                    '').split(
                                            '\n') if line.strip()])
                doc.add_paragraph(clean_text)

                output = io.BytesIO()
                doc.save(output)
                output.seek(0)

                st.download_button(
                    label="📄 导出大纲为 Word",
                    data=output.getvalue(),
                    file_name=f"{base_name}_分集大纲.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )
            except Exception:
                pass

        # 🎯 动态对齐：完美居中排版
        if not is_readonly:
            has_changes = edited_outline != st.session_state.novel_outline

            if has_changes:
                col_save, col_confirm = st.columns(2)
                with col_save:
                    if st.button("💾 保存修改", use_container_width=True):
                        st.session_state.novel_outline = edited_outline
                        st.success("✅ 大纲已保存")
                        st.rerun()
                with col_confirm:
                    if st.button(f"🎬 第二步：确认无误，开始生成 {total_episodes} 集分镜", type="primary",
                                 use_container_width=True):
                        st.session_state.novel_is_generating = True
                        st.rerun()
            else:
                if st.button(f"🎬 第二步：确认无误，开始生成 {total_episodes} 集分镜", type="primary",
                             use_container_width=True):
                    st.session_state.novel_is_generating = True
                    st.rerun()
    st.divider()


def _execute_generation_flow(llm_service, df, total_episodes: int = 20):
    logger.info(f"🚀 开始生成 {total_episodes} 集大纲")
    processor = NovelModeProcessor(llm_service, total_episodes=total_episodes)

    # 🌟 修复：加回大纲的预估时间计算
    estimated_time = "1-2分钟" if total_episodes <= 20 else ("2-3分钟" if total_episodes <= 50 else "3-5分钟")
    progress_bar = st.progress(0, text=f"🚀 正在分析原著，构思 {total_episodes} 集大纲（预计 {estimated_time}）...")
    try:
        with st.spinner():
            outline_text = processor.generate_outline(df, on_progress=lambda msg, val: progress_bar.progress(
                5 + int(val * 0.15), text=msg
            ))
        st.session_state.novel_outline = outline_text
        progress_bar.progress(1.0, text=f"✅ 大纲生成完成")
        st.session_state.novel_is_generating = False
        st.rerun()
    except Exception as e:
        show_inline_error(error_message=str(e), error_type="unknown", show_details=True)
        st.session_state.novel_is_generating = False
        time.sleep(0.5)
        st.rerun()


def _execute_scripts_generation(llm_service, df, total_episodes: int = 20):
    logger.info(f"🎬 开始生成 {total_episodes} 集分镜")
    processor = NovelModeProcessor(llm_service, total_episodes=total_episodes)

    # 🌟 修复：加回分镜多线程的预估时间计算
    min_minutes = max(2, total_episodes // 10)
    max_minutes = max(3, (total_episodes + 5) // 8)
    estimated_minutes = f"{min_minutes}-{max_minutes}"
    progress_bar = st.progress(0,
                               text=f"🚀 多线程引擎启动，准备生成 {total_episodes} 集分镜（预计 {estimated_minutes} 分钟）...")
    try:
        results = processor.process(df, on_progress=lambda msg, val: progress_bar.progress(
            5 + int(val * 0.95), text=msg
        ))
        st.session_state.novel_results = results
        StateManager.set_results(results)
        progress_bar.progress(1.0, text="✅ 全部生成完成")
        st.session_state.novel_is_generating = False
        st.rerun()
    except Exception as e:
        show_inline_error(error_message=str(e), error_type="unknown", show_details=True)
        st.session_state.novel_is_generating = False
        time.sleep(0.5)
        st.rerun()


def _execute_retry_flow(llm_service, df, error_keys, existing_results):
    total_episodes = st.session_state.get('total_episodes', 20)
    processor = NovelModeProcessor(llm_service, total_episodes=total_episodes)
    progress_bar = st.progress(0, text="🔄 正在召唤 AI 填补空缺集数...")

    try:
        results = processor.process(df, on_progress=lambda msg, val: progress_bar.progress(min(val / 100, 1.0),
                                                                                           text=f"⏳ 补全: {val}% | {msg}"),
                                    existing_results=existing_results)
        st.session_state.novel_results = results
        update_count = sum(1 for k in error_keys if results.get(k) is not None and not (
                    isinstance(results.get(k), str) and results.get(k).startswith("❌")) and not (
                    hasattr(results.get(k), 'empty') and results.get(k).empty))
        progress_bar.progress(1.0, text=f"✅ 补全完成，成功修复 {update_count}/{len(error_keys)} 集")
        st.session_state.novel_is_generating = False
        st.rerun()
    except Exception as e:
        show_inline_error(error_message=str(e), error_type="unknown", show_details=True)
        st.session_state.novel_is_generating = False
        time.sleep(0.5)
        st.rerun()


def _sort_and_clean_results(results: dict) -> dict:
    try:
        def extract_num(key):
            match = re.search(r'\d+', str(key))
            return int(match.group()) if match else 0

        return {k: results[k] for k in sorted(results.keys(), key=extract_num)}
    except:
        return results