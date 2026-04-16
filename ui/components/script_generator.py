"""
剧本生成组件模块 - 已集成进度条、流式输出与默认《低画质人生》创意
"""
import io
import logging
import re
import time
import zipfile

import pandas as pd
import streamlit as st
from core.processor import NovelModeProcessor
from core.prompts import PromptTemplates

logger = logging.getLogger("ScriptMaster.ScriptGenerator")


def export_outline_to_word(outline_text: str, total_eps: int = 30) -> bytes:
    """导出大纲为 Word 文档"""
    from docx import Document
    doc = Document()
    doc.add_heading(f'{total_eps}集分集大纲', 0)
    doc.add_paragraph(outline_text)
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


def init_script_state():
    """初始化剧本生成相关的session state"""
    if 'generated_acts' not in st.session_state:
        st.session_state.generated_acts = None
    if 'selected_act' not in st.session_state:
        st.session_state.selected_act = None
    if 'outline' not in st.session_state:
        st.session_state.outline = None
    if 'scripts' not in st.session_state:
        st.session_state.scripts = {}
    if 'script_is_generating' not in st.session_state:
        st.session_state.script_is_generating = False
    if 'script_total_episodes' not in st.session_state:
        st.session_state.script_total_episodes = 30


def extract_csv_from_text(text):
    """提取包含表头的 CSV 数据块"""
    if not text: return ""
    text = re.sub(r'```csv\s*', '', text)
    text = re.sub(r'```\s*', '', text)
    pattern = r"镜号,场景,画面内容 \(Visual\),台词 \(Dialogue\) & 音效 \(SFX\).*"
    match = re.search(pattern, text, re.DOTALL)
    if match:
        return match.group(0).strip()
    return text.strip()


def _render_step1_buttons(api_ready: bool):
    if not api_ready:
        st.warning("⚠️ 请在侧边栏验证 API 配置")
        st.button("🚀 开始创作三幕式", disabled=True, key="gen_act_btn_dis")
        return "waiting_api"
    elif not st.session_state.generated_acts:
        if st.button("🚀 开始创作三幕式", type="primary", key="gen_act_btn",
                     disabled=st.session_state.get('script_is_generating', False)):
            st.session_state.script_is_generating = True
            st.rerun()
        return "ready_to_generate"
    else:
        if st.button("🗑️ 清除三幕式并重新生成", width='stretch',
                     disabled=st.session_state.get('script_is_generating', False)):
            st.session_state.generated_acts = None
            st.session_state.selected_act = None
            st.rerun()
        return "already_generated"


def _execute_step1_generation(llm_service, original_idea: str):
    full_response = ""
    prompt = PromptTemplates.ACT_GEN_TASK.format(original_idea=original_idea)
    placeholder = st.empty()

    idea_length = len(original_idea)
    if idea_length < 100:
        estimated_time = "30-60秒"
    elif idea_length < 500:
        estimated_time = "1-2分钟"
    else:
        estimated_time = "2-3分钟"

    try:
        with st.spinner(f"🤖 AI 编剧正在构思剧情（预计{estimated_time}）..."):
            for chunk in llm_service.generate_stream("你是一个专业编剧，擅长爆款短剧设定。", prompt):
                full_response += chunk
                placeholder.markdown(full_response + "▌")

        placeholder.markdown(full_response)

        if full_response.startswith("❌"):
            st.error("❌ 生成失败：请检查 API 配置和网络连接")
            st.session_state.script_is_generating = False
        else:
            st.session_state.generated_acts = full_response
            st.success("✅ 三幕式构架生成完成！")
            st.session_state.script_is_generating = False
            st.rerun()
    except Exception as e:
        st.error(f"❌ 生成异常: {str(e)[:60]}")
        st.session_state.script_is_generating = False
        st.rerun()


def _render_step1_selection():
    st.divider()
    st.markdown("**📋 选择并修改三幕式构架**")

    acts_list = st.session_state.generated_acts.split("---")
    acts_list = [a.strip() for a in acts_list if a.strip()]

    if len(acts_list) > 1:
        act_titles = [a.split("\n")[0][:30] + "..." for a in acts_list]
        if 'selected_act_index' not in st.session_state:
            st.session_state.selected_act_index = 0

        st.session_state.selected_act_index = st.radio(
            "请选择要生成大纲的创意：",
            options=range(len(acts_list)),
            format_func=lambda x: f"创意 {x + 1}: {act_titles[x]}",
            horizontal=True,
            key="act_radio"
        )
        return acts_list[st.session_state.selected_act_index], acts_list
    else:
        st.session_state.selected_act_index = 0
        return st.session_state.generated_acts, acts_list


def _render_step1_editor(current_act: str, acts_list: list):
    # 🌟 优化：如果已经到了大纲生成阶段，第一步自动折叠
    is_expanded = not bool(st.session_state.get('outline'))
    with st.expander("✏️ 点击编辑选中的三幕式构架", expanded=is_expanded):
        edited_act = st.text_area(
            "构架内容（可直接修改）：",
            value=current_act,
            height=300,
            label_visibility="collapsed",
            disabled=st.session_state.get('script_is_generating', False)
        )

        if st.button("💾 保存修改", width='stretch'):
            if len(acts_list) > 1:
                acts_list[st.session_state.selected_act_index] = edited_act
                st.session_state.generated_acts = "\n\n---\n\n".join(acts_list)
            else:
                st.session_state.generated_acts = edited_act

            st.session_state.selected_act = edited_act
            st.success("✅ 创意已保存并同步")
            st.rerun()


def render_step_1_acts(llm_service):
    try:
        st.markdown('<div class="step-header">第一步：核心创意与三幕式构架</div>', unsafe_allow_html=True)

        default_idea = """Resolution: Low《低画质人生》\n2099年，视觉感知成为一种昂贵的订阅服务。富人享受着8K HDR的极致世界，而像凯这样的穷人只能活在"经济模式"里--一个模糊、像素化的144p噩梦。凯为了看清病危母亲的脸，在黑市购买了违禁芯片，结果发现"高清"世界里，统治者其实是食人怪物，而"低画质"只是为了掩盖真相的滤镜。"""

        original_idea = st.text_area("请输入原始创意：", value=default_idea, height=150,
                                     disabled=st.session_state.get('script_is_generating', False))

        api_ready = st.session_state.get("api_validated", False)

        _render_step1_buttons(api_ready)

        if st.session_state.script_is_generating and not st.session_state.generated_acts:
            _execute_step1_generation(llm_service, original_idea)

        if st.session_state.generated_acts and not st.session_state.generated_acts.startswith("❌"):
            current_act, acts_list = _render_step1_selection()
            _render_step1_editor(current_act, acts_list)

    except Exception as e:
        st.error(f"❌ 第一步加载异常: {str(e)[:60]}")
        logger.error(f"❌ [ScriptStep1] 渲染异常: {str(e)}", exc_info=True)
        st.session_state.script_is_generating = False


def _render_step2_buttons(api_ready: bool, total_episodes: int, batch_count: int):
    st.markdown("**⚙️ 大纲生成配置**")
    col1, col2 = st.columns([3, 1])
    with col1:
        total_episodes = st.number_input(
            "总集数", min_value=10, max_value=100,
            value=st.session_state.get('script_total_episodes', 30), step=5
        )
        st.session_state.script_total_episodes = total_episodes
    with col2:
        st.markdown(
            f"""<div style="background-color: #e8f4fd; padding: 10px; border-radius: 5px; text-align: center;">
                <p style="margin: 0; font-size: 14px;">📦 分 <b>{batch_count} 次</b> 生成<br>
                <span style="color: #666; font-size: 12px;">共{total_episodes}集</span></p>
            </div>""", unsafe_allow_html=True
        )

    if not api_ready:
        st.warning("⚠️ 请在侧边栏验证 API 配置")
        st.button(f"📈 生成 {total_episodes} 集大纲", disabled=True, width='stretch')
        return "waiting_api"
    else:
        if st.button(f"📈 生成 {total_episodes} 集大纲", type="primary",
                     disabled=st.session_state.get('script_is_generating', False), width='stretch'):
            st.session_state.script_is_generating = True
            st.rerun()
        return "ready_to_generate"


def _execute_step2_generation(llm_service, source_act: str):
    full_response = ""
    source_length = len(source_act)
    total_episodes = st.session_state.get('script_total_episodes', 30)

    if total_episodes <= 20:
        estimated_time = "1-2分钟" if source_length < 500 else "2-3分钟"
    elif total_episodes <= 50:
        estimated_time = "2-3分钟" if source_length < 500 else "3-5分钟"
    else:
        estimated_time = "3-5分钟" if source_length < 500 else "5-8分钟"

    prompt = PromptTemplates.OUTLINE_TASK.format(user_choice=source_act, total_episodes=total_episodes)
    placeholder = st.empty()

    try:
        with st.spinner(f"📖 AI 正在扩展大纲（预计{estimated_time}）..."):
            for chunk in llm_service.generate_stream(PromptTemplates.OUTLINE_SYSTEM, prompt):
                full_response += chunk
                placeholder.markdown(full_response + "▌")

        if full_response.startswith("❌"):
            st.error("❌ 大纲生成失败：请检查 API 配置和网络连接")
            st.session_state.script_is_generating = False
        else:
            st.session_state.outline = full_response
            st.success(f"✅ {total_episodes} 集大纲生成完成！")
            st.session_state.script_is_generating = False
            st.rerun()
    except Exception as e:
        st.error(f"❌ 大纲生成异常: {str(e)[:60]}")
        st.session_state.script_is_generating = False
        st.rerun()


def _extract_script_title(act_text: str) -> str:
    if not act_text: return "未命名剧本"
    first_line = act_text.split('\n')[0].strip()
    match = re.search(r'《(.+?)》', first_line)
    if match: return match.group(1)
    return first_line[:20] if first_line else "未命名剧本"


def _render_step2_editor_and_export():
    st.divider()
    st.markdown("**📖 分集大纲预览**")

    # 🌟 优化：如果第三步生成了最终分镜，则大纲自动折叠，高度减小
    is_expanded = not bool(st.session_state.get('scripts'))
    with st.expander("📋 点击查看/编辑大纲", expanded=is_expanded):
        edited_outline = st.text_area(
            "大纲内容（可直接修改）：",
            value=st.session_state.outline,
            height=300 if not is_expanded else 450,
            label_visibility="collapsed",
            disabled=st.session_state.get('script_is_generating', False)
        )

        if edited_outline != st.session_state.outline:
            if st.button("💾 保存修改", width='stretch'):
                st.session_state.outline = edited_outline
                st.success("✅ 大纲已保存")
                st.rerun()

        total_eps = st.session_state.get('script_total_episodes', 30)
        word_bytes = export_outline_to_word(st.session_state.outline, total_eps)

        source_act = st.session_state.get('selected_act') or st.session_state.get('generated_acts', '')
        script_title = _extract_script_title(source_act)

        # 🚨 净化文件名，防止非法字符报错
        safe_title = re.sub(r'[\\/*?:"<>|]', "_", script_title)

        st.download_button(
            label="📄 导出为 Word",
            data=word_bytes,
            file_name=f"{safe_title}_{total_eps}集大纲.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            width='stretch'
        )
        st.info("💡 提示: 大纲确认无误后，请向下滚动到第三步，点击「开始批量执行」生成分镜")


def render_step_2_outline(llm_service):
    try:
        if not st.session_state.get('selected_act') and not st.session_state.get('generated_acts'):
            return

        st.markdown('<div class="step-header">第二步：扩展分集大纲</div>', unsafe_allow_html=True)

        source_act = st.session_state.get('selected_act') or st.session_state.get('generated_acts')
        total_episodes = st.session_state.get('script_total_episodes', 30)
        batch_count = (total_episodes + 2) // 3

        if not st.session_state.outline:
            _render_step2_buttons(st.session_state.get("api_validated", False), total_episodes, batch_count)
        else:
            st.markdown("**⚙️ 大纲管理**")
            if st.button("🗑️ 清除大纲并重新生成", width='stretch',
                         disabled=st.session_state.get('script_is_generating', False)):
                st.session_state.outline = None
                st.rerun()

        if st.session_state.script_is_generating and not st.session_state.outline:
            _execute_step2_generation(llm_service, source_act)
            return

        if st.session_state.outline and not st.session_state.outline.startswith("❌"):
            _render_step2_editor_and_export()

    except Exception as e:
        st.error(f"❌ 第二步加载异常: {str(e)[:60]}")
        st.session_state.script_is_generating = False


def _execute_step3_generation(llm_service, total_episodes: int):
    min_minutes = max(2, total_episodes // 10)
    max_minutes = max(3, (total_episodes + 5) // 8)
    estimated_minutes = f"{min_minutes}-{max_minutes}"
    progress_bar = st.progress(0, text=f"🚀 并行生成 {total_episodes} 集分镜（预计{estimated_minutes}分钟）...")
    try:
        # 🚨 核心修复：每次运行动态实例化 Processor，绝不使用 session 里的死缓存，确保集数实时更新！
        processor = NovelModeProcessor(llm_service, total_episodes=total_episodes)
        processor.user_template = PromptTemplates.BATCH_SCRIPT_PROMPT
        processor.system_prompt = PromptTemplates.SCRIPT_SYSTEM

        res = processor.process(
            full_text=st.session_state.outline,
            on_progress=lambda msg, val: progress_bar.progress(val / 100, text=msg)
        )
        st.session_state.scripts = {k: v.to_csv(index=False) if hasattr(v, 'to_csv') else v for k, v in res.items()}
        progress_bar.progress(1.0, text="✅ 生成完成")
        st.session_state.script_is_generating = False
        time.sleep(0.5)
        st.rerun()
    except Exception as e:
        st.error(f"❌ 生成失败: {str(e)}")
        st.session_state.script_is_generating = False
        time.sleep(0.5)
        st.rerun()


def _execute_step3_retry(llm_service, results: dict, total_episodes: int):
    progress_bar = st.progress(0, text="🔄 正在召唤 AI 填补空缺集数...")
    try:
        processor = NovelModeProcessor(llm_service, total_episodes=total_episodes)
        processor.user_template = PromptTemplates.BATCH_SCRIPT_PROMPT
        processor.system_prompt = PromptTemplates.SCRIPT_SYSTEM

        res = processor.process(
            full_text=st.session_state.outline,
            existing_results={k: (pd.read_csv(io.StringIO(v)) if v and not v.startswith("❌") else v) for k, v in
                              results.items()},
            on_progress=lambda msg, val: progress_bar.progress(val / 100, text=msg)
        )
        st.session_state.scripts = {k: v.to_csv(index=False) if hasattr(v, 'to_csv') else v for k, v in res.items()}
        progress_bar.progress(1.0, text="✅ 补全完成")
        st.session_state.script_is_generating = False
        time.sleep(0.5)
        st.rerun()
    except Exception as e:
        st.error(f"❌ 补全失败: {str(e)}")
        st.session_state.script_is_generating = False
        time.sleep(0.5)
        st.rerun()


def render_step_3_scripts(llm_service):
    try:
        if not st.session_state.outline:
            return

        st.markdown('<div class="step-header">第三步：批量生成详细分镜</div>', unsafe_allow_html=True)
        total_episodes = st.session_state.get('script_total_episodes', 30)
        results = st.session_state.scripts

        if st.session_state.script_is_generating:
            st.info("⏳ 正在批量生成分镜，请稍候...")
        elif not results:
            api_ready = st.session_state.get("api_validated", False)
            if not api_ready:
                st.warning("⚠️ 请在侧边栏验证 API 配置")
            else:
                if st.button("🔥 开始批量执行", type="primary",
                             disabled=st.session_state.get('script_is_generating', False), width='stretch'):
                    st.session_state.script_is_generating = True
                    st.rerun()
        else:
            if st.button("🗑️ 清除当前结果并重新开始", width='stretch',
                         disabled=st.session_state.get('script_is_generating', False)):
                st.session_state.scripts = {}
                st.rerun()

        if st.session_state.script_is_generating and not st.session_state.scripts:
            _execute_step3_generation(llm_service, total_episodes)
            return

        if st.session_state.script_is_generating and st.session_state.scripts:
            error_keys = [k for k, v in st.session_state.scripts.items() if isinstance(v, str) and v.startswith("❌")]
            if error_keys:
                _execute_step3_retry(llm_service, st.session_state.scripts, total_episodes)
                return
        elif st.session_state.scripts:
            error_keys = [k for k, v in st.session_state.scripts.items() if isinstance(v, str) and v.startswith("❌")]
            if error_keys and not st.session_state.get('script_is_generating', False):
                retry_label = f"🔄 补全缺失集数 ({len(error_keys)}集)"
                if st.button(retry_label, width='stretch'):
                    st.session_state.script_is_generating = True
                    st.rerun()

    except Exception as e:
        st.error(f"❌ 第三步加载异常: {str(e)[:60]}")
        st.session_state.script_is_generating = False


def display_final_tables():
    if not st.session_state.scripts:
        return

    st.divider()
    st.markdown("### 🎬 剧本库预览")
    results = st.session_state.scripts
    sorted_keys = sorted(results.keys(), key=lambda x: int(re.findall(r'\d+', x)[0]))

    total_shots, valid_eps = 0, 0
    for k in sorted_keys:
        if results[k] and not results[k].startswith("❌"):
            try:
                df = pd.read_csv(io.StringIO(results[k]))
                total_shots += len(df)
                valid_eps += 1
            except:
                pass

    st.markdown(f"📊 **总计：{valid_eps} 集，{total_shots} 个镜头**")

    low_shot_eps = []
    for k in sorted_keys:
        if results[k] and not results[k].startswith("❌"):
            try:
                shot_count = pd.read_csv(io.StringIO(results[k])).shape[0]
                if shot_count < 20:
                    low_shot_eps.append((k, shot_count))
            except:
                pass

    if low_shot_eps:
        st.warning(f"⚠️ 以下 {len(low_shot_eps)} 集镜头数不足 20 个，建议点击「补全缺失集数」重试：")
        for ep, count in low_shot_eps:
            st.markdown(f"- **{ep}**: {count} 个镜头")

    selected_key = st.selectbox("🎬 选择查看集数：", options=sorted_keys, format_func=lambda x: f"📺 {x}", index=0)

    if selected_key:
        content = results[selected_key]
        try:
            df = pd.read_csv(io.StringIO(content))
            st.dataframe(df, width='stretch', hide_index=True)
            csv_bytes = df.to_csv(index=False).encode('utf-8-sig')
            st.download_button(f"📥 下载{selected_key}", csv_bytes, f"{selected_key}.csv", "text/csv",
                               key=f"btn_{selected_key}")
        except:
            st.text_area("内容解析异常，请检查：", value=content, height=200)

    _render_batch_download_script(results)


def _render_batch_download_script(results: dict):
    st.divider()
    st.markdown("**📥 批量下载**")

    source_act = st.session_state.get('selected_act') or st.session_state.get('generated_acts', '')
    match = re.search(r'《(.+?)》', source_act)
    script_title = match.group(1) if match else "未命名剧本"

    # 🚨 净化文件名，防止因特殊符号导致 ZIP 打包崩溃
    safe_title = re.sub(r'[\\/*?:"<>|]', "_", script_title)

    col1, col2 = st.columns(2)
    with col1:
        excel_output = io.BytesIO()
        with pd.ExcelWriter(excel_output, engine='openpyxl') as writer:
            for k, v in results.items():
                if v and not v.startswith("❌"):
                    try:
                        pd.read_csv(io.StringIO(v)).to_excel(writer, sheet_name=k, index=False)
                    except:
                        pass
        excel_output.seek(0)
        st.download_button("📊 下载 Excel（多工作表）", excel_output.getvalue(),
                           f"{safe_title}_剧本分镜.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           width='stretch')

    with col2:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
            for t, c in results.items():
                if c and not c.startswith("❌"):
                    try:
                        z.writestr(f"{t}_分镜.csv", c.encode('utf-8-sig'))
                    except:
                        pass
        buf.seek(0)

        st.download_button(
            label="📦 点击下载 ZIP（分集 CSV）",
            data=buf.getvalue(),
            file_name=f"{safe_title}_分镜集数包.zip",
            mime="application/zip",
            width='stretch',
            key="btn_zip_script"
        )


def render_script_generation_mode(llm_service):
    st.markdown("""
        <style>
            .stTextArea textarea { color: #000000 !important; }
            .step-header { padding: 10px; background: #f0f2f6; border-left: 5px solid #E3700D; font-weight: bold; margin: 15px 0; color: #2C2A29;}
        </style>
    """, unsafe_allow_html=True)

    init_script_state()
    render_step_1_acts(llm_service)
    render_step_2_outline(llm_service)
    render_step_3_scripts(llm_service)
    display_final_tables()